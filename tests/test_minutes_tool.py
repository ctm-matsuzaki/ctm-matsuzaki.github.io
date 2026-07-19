from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from minutes_tool.excel_writer import MinutesExcelWriter, normalize_basic_datetime, normalize_basic_place
from minutes_tool.filename import build_display_meeting_name, build_output_filename, build_sheet_name
from minutes_tool.parser import MinutesParser


TEMPLATE = Path("minutes_tool/resources/会議議事録テンプレのコピー.xlsx")


def _write(text: str, tmp_path: Path):
    minutes = MinutesParser().parse(text)
    output = tmp_path / build_output_filename(minutes.meeting_name, minutes.date_time, short_meeting_name=minutes.short_meeting_name)
    MinutesExcelWriter(TEMPLATE).write(minutes, output)
    return minutes, output, load_workbook(output).active


def _base_text(meeting_name: str = "燈心会人員配置会議", date: str = "2026年7月16日") -> str:
    return f"""
# 基本情報
**会議名**
{meeting_name}
**日時**
{date} 14:00～15:00
**場所**
オンライン
**出席者**
松崎さん、名倉さん

# 全体要約
全体要約の本文です。

# 議題① 人員配置
## 内容
人員配置について確認しました。
## 決定事項
看護師1名を追加で調整します。
受付体制を見直します。
## 未決事項・継続検討事項
採用時期は継続検討です。

# ToDo
| 担当 | 期限 | 内容 |
|---|---|---|
| 松崎さん | 2026年7月20日 | 資料を共有する |

# 次回会議
日時：2026年7月30日（木）14:00～15:00　場所：オンライン
"""


def test_meeting_date_and_short_name_within_15_chars(tmp_path: Path) -> None:
    minutes, output, sheet = _write(_base_text(), tmp_path)

    assert minutes.meeting_name == "燈心会人員配置会議"
    assert output.name == "7.16_燈心会人員配置会議.xlsx"
    assert sheet.title == "7.16"
    assert sheet["C4"].value == "燈心会人員配置会議"
    assert len(sheet["C4"].value) <= 15


def test_long_meeting_name_is_shortened(tmp_path: Path) -> None:
    text = _base_text("燈心会の人員配置についての打ち合わせに関する会議")
    _, output, sheet = _write(text, tmp_path)

    assert len(sheet["C4"].value) <= 15
    assert output.name.startswith("7.16_")
    assert len(output.stem.split("_", 1)[1]) <= 15


def test_short_meeting_name_in_markdown_has_priority(tmp_path: Path) -> None:
    text = _base_text("燈心会の人員配置についての打ち合わせに関する会議").replace(
        "**日時**",
        "**短い会議名**\n燈心会人員配置会議\n**日時**",
    )
    _, output, sheet = _write(text, tmp_path)

    assert sheet["C4"].value == "燈心会人員配置会議"
    assert output.name == "7.16_燈心会人員配置会議.xlsx"


def test_missing_meeting_date_uses_jst_today() -> None:
    assert build_output_filename("会議", "", now=datetime(2026, 7, 16, 23, 30)) == "7.16_会議.xlsx"
    assert build_sheet_name("", now=datetime(2026, 7, 16, 23, 30)) == "7.16"


def test_basic_datetime_outputs_only_date_and_time(tmp_path: Path) -> None:
    text = _base_text(date="2026年7月17日（金） 10:00〜11:00（オンライン開催）")
    _, _, sheet = _write(text, tmp_path)

    assert sheet["C5"].value == "2026年7月17日 10:00〜11:00"


def test_basic_datetime_outputs_date_only_when_time_is_missing(tmp_path: Path) -> None:
    text = _base_text().replace("2026年7月16日 14:00～15:00", "2026年7月17日（金）（補足：時間未定）")
    _, _, sheet = _write(text, tmp_path)

    assert sheet["C5"].value == "2026年7月17日"


def test_basic_datetime_uses_today_when_date_is_missing() -> None:
    assert normalize_basic_datetime("", now=datetime(2026, 7, 17, 10, 0)) == "2026年7月17日"


def test_basic_place_is_simplified_for_online_and_physical_locations(tmp_path: Path) -> None:
    online_text = _base_text().replace("オンライン", "Zoom（URL：https://example.com、接続方法は別途共有）", 1)
    _, _, online_sheet = _write(online_text, tmp_path)
    assert online_sheet["C6"].value == "オンライン"

    physical_text = _base_text().replace("オンライン", "燈心会 3階会議室（大阪府大阪市、受付で入館）", 1)
    _, _, physical_sheet = _write(physical_text, tmp_path)
    assert physical_sheet["C6"].value == "燈心会 3階会議室"


def test_basic_place_unknown_becomes_blank() -> None:
    assert normalize_basic_place("不明（確認中）") == ""


def test_next_meeting_has_date_time_and_place(tmp_path: Path) -> None:
    _, _, sheet = _write(_base_text(), tmp_path)

    next_value = _cell_text(sheet, "日時：2026年7月30日")
    assert "場所：オンライン" in next_value


def test_next_meeting_has_date_time_only(tmp_path: Path) -> None:
    text = _base_text().replace("日時：2026年7月30日（木）14:00～15:00　場所：オンライン", "日時：2026年7月30日（木）14:00～15:00")
    _, _, sheet = _write(text, tmp_path)

    next_value = _cell_text(sheet, "日時：2026年7月30日")
    assert "場所：" not in next_value


def test_empty_next_meeting_is_blank_and_highlight_target(tmp_path: Path) -> None:
    text = _base_text().replace("# 次回会議\n日時：2026年7月30日（木）14:00～15:00　場所：オンライン", "# 次回会議\n")
    _, _, sheet = _write(text, tmp_path)

    next_row = _row_containing(sheet, "【次回日程】") + 1
    assert sheet.cell(next_row, 2).value in (None, "")
    assert any(str(rule.sqref) == f"B{next_row}:E{next_row}" for rule in sheet.conditional_formatting)


def test_multiple_decisions_are_output_above_meeting_content(tmp_path: Path) -> None:
    _, _, sheet = _write(_base_text(), tmp_path)

    decision_row = _row_containing(sheet, "看護師1名を追加")
    content_row = _row_containing(sheet, "人員配置について確認")
    assert decision_row < content_row
    assert "受付体制を見直します" in sheet.cell(decision_row + 1, 3).value


def test_long_content_and_summary_have_wrap_and_height(tmp_path: Path) -> None:
    long_sentence = "長文です。" * 80
    long_content = f"打合せ内容の{long_sentence}"
    text = _base_text().replace("全体要約の本文です。", long_sentence).replace("人員配置について確認しました。", long_content)
    _, _, sheet = _write(text, tmp_path)

    assert sheet["B10"].alignment.horizontal == "left"
    assert sheet["B10"].alignment.vertical == "top"
    assert sheet["B10"].alignment.wrap_text is True
    assert sheet.row_dimensions[10].height > 22
    content_row = _row_containing(sheet, "打合せ内容の")
    assert sheet.cell(content_row, 3).alignment.horizontal == "left"
    assert sheet.row_dimensions[content_row].height > 22


def test_invalid_windows_filename_chars_are_removed() -> None:
    assert build_display_meeting_name(r'会議\\/:*?"<>|名') == "会議名"
    assert build_output_filename(r'会議\\/:*?"<>|名', "2026年7月16日") == "7.16_会議名.xlsx"


def test_file_date_and_sheet_date_are_the_same(tmp_path: Path) -> None:
    _, output, sheet = _write(_base_text(), tmp_path)

    assert output.stem.split("_", 1)[0] == sheet.title


def test_template_can_be_loaded_after_write(tmp_path: Path) -> None:
    _, output, _ = _write(_base_text(), tmp_path)

    workbook = load_workbook(output)
    assert workbook.active.title == "7.16"
    assert len(workbook.active.tables) == 0

    with ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "xl/tables/table1.xml" not in names
        assert not [name for name in names if name.startswith("xl/tables/")]
        for name in names:
            if name.startswith("xl/worksheets/_rels/"):
                rels = archive.read(name).decode("utf-8")
                assert "relationships/table" not in rels
                assert "table1.xml" not in rels
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
        assert "table+xml" not in content_types


def test_dynamic_borders_with_one_agenda_and_one_todo(tmp_path: Path) -> None:
    _, output, sheet = _write(_dynamic_text(agenda_count=1, todo_count=1), tmp_path)

    _assert_dynamic_block_borders(sheet, agenda_count=1, todo_count=1)
    _assert_todo_outer_border(sheet, todo_count=1)
    _assert_no_table_parts(output)


def test_dynamic_borders_with_three_agendas_three_todos_and_pending(tmp_path: Path) -> None:
    _, output, sheet = _write(_dynamic_text(agenda_count=3, todo_count=3, with_pending=True), tmp_path)

    _assert_dynamic_block_borders(sheet, agenda_count=3, todo_count=3)
    _assert_todo_outer_border(sheet, todo_count=3)
    assert _row_containing(sheet, "未決事項・継続検討事項")
    _assert_no_table_parts(output)


def test_dynamic_borders_with_five_agendas_and_no_todos(tmp_path: Path) -> None:
    _, output, sheet = _write(_dynamic_text(agenda_count=5, todo_count=0, with_pending=False), tmp_path)

    _assert_dynamic_block_borders(sheet, agenda_count=5, todo_count=1)
    _assert_todo_outer_border(sheet, todo_count=1)
    _assert_no_table_parts(output)


def _dynamic_text(agenda_count: int, todo_count: int, with_pending: bool = True) -> str:
    agenda_parts = []
    for index in range(1, agenda_count + 1):
        pending = "未決事項は継続確認です。" if with_pending else "該当なし"
        decision = "該当なし" if index == 1 else f"決定事項{index}を実施します。\n追加決定{index}を確認しました。"
        content = "短い内容です。" if index == 1 else f"打合せ内容{index}です。\n複数行の内容{index}です。"
        agenda_parts.append(
            f"""
# 議題{index} 議題タイトル{index}
## 内容
{content}
## 決定事項
{decision}
## 未決事項・継続検討事項
{pending}
"""
        )

    if todo_count:
        todo_rows = "\n".join(
            f"| 担当{index} | 2026年7月{20 + index}日 | タスク{index}を対応する |"
            for index in range(1, todo_count + 1)
        )
        todo_section = f"""
# ToDo
| 担当 | 期限 | 内容 |
|---|---|---|
{todo_rows}
"""
    else:
        todo_section = """
# ToDo
該当なし
"""

    return f"""
# 基本情報
**会議名**
罫線確認会議
**日時**
2026年7月16日 14:00～15:00
**場所**
オンライン
**出席者**
松崎さん、名倉さん

# 全体要約
全体要約の本文です。

{''.join(agenda_parts)}
{todo_section}
# 次回会議
日時：2026年7月30日（木）14:00～15:00　場所：オンライン
"""


def _assert_dynamic_block_borders(sheet, agenda_count: int, todo_count: int) -> None:
    decision_start = _row_containing(sheet, "決定事項") + 1
    content_start = _row_containing(sheet, "打合せ内容") + 1
    todo_start = _row_containing(sheet, "ToDoタスク") + 1

    for index in range(agenda_count):
        decision_end = decision_start + ((index + 1) * 5) - 1
        content_end = content_start + ((index + 1) * 5) - 1
        _assert_bottom_styles(sheet, decision_end, range(2, 6), "medium")
        _assert_bottom_styles(sheet, content_end, range(2, 6), "medium")
        assert sheet.cell(decision_end, 2).border.left.style == "medium"
        assert sheet.cell(decision_end, 5).border.right.style == "medium"
        assert sheet.cell(content_end, 2).border.left.style == "medium"
        assert sheet.cell(content_end, 5).border.right.style == "medium"

    _assert_bottom_styles(sheet, todo_start + todo_count - 1, range(2, 6), "thin")

    assert sheet.cell(decision_start, 3).border.right.style == "medium"
    assert sheet.cell(content_start, 5).border.right.style == "medium"
    if agenda_count:
        assert sheet.cell(content_start + 1, 3).border.bottom.style == "dotted"
    if todo_count > 1:
        assert sheet.cell(todo_start, 3).border.bottom.style == "dotted"
    assert any(str(merged_range).startswith(f"C{decision_start}:E{decision_start}") for merged_range in sheet.merged_cells.ranges)
    assert any(str(merged_range).startswith(f"C{content_start}:D{content_start}") for merged_range in sheet.merged_cells.ranges)
    assert len(sheet.conditional_formatting) > 0


def _assert_todo_outer_border(sheet, todo_count: int) -> None:
    header_row = _row_containing(sheet, "ToDoタスク")
    first_body_row = header_row + 1
    end_row = first_body_row + todo_count - 1

    for column in range(2, 6):
        assert sheet.cell(header_row, column).border.top.style == "medium"
        assert sheet.cell(end_row, column).border.bottom.style == "thin"

    for row in range(header_row, end_row + 1):
        assert sheet.cell(row, 2).border.left.style == "thin"
        assert sheet.cell(row, 5).border.right.style == "thin"

    assert sheet.cell(header_row, 2).border.right.style == "thin"
    assert sheet.cell(header_row, 3).border.left.style == "thin"
    assert sheet.cell(header_row, 3).border.right.style == "thin"
    assert sheet.cell(header_row, 4).border.right.style == "thin"
    if todo_count > 1:
        assert sheet.cell(first_body_row, 3).border.bottom.style == "dotted"
        assert sheet.cell(first_body_row + 1, 3).border.bottom.style == "dotted"

    spacer_row = end_row + 1
    assert sheet.cell(spacer_row, 2).border.left.style is None
    assert sheet.cell(spacer_row, 5).border.right.style is None


def _assert_bottom_styles(sheet, row: int, columns, expected: str) -> None:
    for column in columns:
        assert sheet.cell(row=row, column=column).border.bottom.style == expected


def _assert_no_table_parts(output: Path) -> None:
    workbook = load_workbook(output)
    assert len(workbook.active.tables) == 0
    with ZipFile(output) as archive:
        names = set(archive.namelist())
        assert not [name for name in names if name.startswith("xl/tables/")]
        for name in names:
            if name.startswith("xl/worksheets/_rels/"):
                rels = archive.read(name).decode("utf-8")
                assert "relationships/table" not in rels
                assert "table1.xml" not in rels


def _row_containing(sheet, text: str) -> int:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and text in str(cell.value):
                return cell.row
    raise AssertionError(f"not found: {text}")


def _cell_text(sheet, text: str) -> str:
    row = _row_containing(sheet, text)
    for cell in sheet[row]:
        if cell.value and text in str(cell.value):
            return str(cell.value)
    raise AssertionError(f"not found: {text}")
