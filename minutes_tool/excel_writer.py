from __future__ import annotations

from copy import copy
from datetime import datetime, timedelta, timezone
from math import ceil
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment
from openpyxl.styles import PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from .filename import build_display_meeting_name, build_sheet_name
from .models import AgendaItem, MeetingMinutes, TodoItem


JST = timezone(timedelta(hours=9))


class MinutesExcelError(RuntimeError):
    """User-facing Excel generation error."""


class MinutesExcelWriter:
    """Write parsed minutes into the fixed CTM Excel template."""

    DYNAMIC_START_ROW = 14
    DYNAMIC_END_ROW = 64
    AGENDA_DECISION_ROWS = 5
    AGENDA_CONTENT_ROWS = 5
    TODO_MIN_ROWS = 1

    SUMMARY_HEADER_ROW = 9
    SUMMARY_BODY_ROW = 10

    def __init__(self, template_path: str | Path) -> None:
        self.template_path = Path(template_path)
        self._missing_ranges: list[str] = []

    def write(self, minutes: MeetingMinutes, output_path: str | Path) -> Path:
        if not self.template_path.exists():
            raise FileNotFoundError(f"テンプレートが見つかりません: {self.template_path}")

        output = Path(output_path)
        if output.suffix.lower() != ".xlsx":
            output = output.with_suffix(".xlsx")
        output.parent.mkdir(parents=True, exist_ok=True)

        try:
            workbook = load_workbook(self.template_path)
        except InvalidFileException as exc:
            raise MinutesExcelError(
                "選択されたテンプレートをExcelファイルとして読み込めませんでした。\n"
                "拡張子が .xlsx の会議議事録テンプレートを選択してください。"
            ) from exc
        except PermissionError:
            raise
        except Exception as exc:
            raise MinutesExcelError(
                "Excelテンプレートを読み込めませんでした。\n"
                "テンプレートファイルが破損していないか確認してください。"
            ) from exc

        sheet = workbook.active
        self._validate_template(sheet)
        sheet.title = build_sheet_name(minutes.date_time)

        self._write_basic_info(sheet, minutes)
        self._write_summary(sheet, minutes)
        self._rebuild_dynamic_area(sheet, minutes)
        self._format_for_readability(sheet)
        self._apply_missing_value_highlights(sheet)
        self._remove_excel_tables(workbook)

        try:
            workbook.save(output)
        except PermissionError:
            raise
        except Exception as exc:
            raise MinutesExcelError(
                "Excelファイルの保存中にエラーが発生しました。\n"
                "保存先、ファイル名、Excelで同名ファイルを開いていないかを確認してください。"
            ) from exc
        return output

    def _remove_excel_tables(self, workbook) -> None:
        for sheet in workbook.worksheets:
            for table_name in list(sheet.tables.keys()):
                del sheet.tables[table_name]

    def _validate_template(self, sheet: Worksheet) -> None:
        required_values = {
            "B2": "会議議事録",
            "B4": "会議名",
            "B5": "日時",
            "B6": "場所",
            "B7": "出席者",
            "B9": "全体の要約",
            "B14": "議題",
            "C14": "決定事項",
            "B31": "議題",
            "C31": "打合せ内容",
            "E31": "発言者",
            "B48": "議題",
            "C48": "ToDoタスク",
            "D48": "担当者",
            "E48": "期限",
        }
        missing = [cell for cell, value in required_values.items() if sheet[cell].value != value]
        if missing:
            raise MinutesExcelError(
                "選択されたExcelは、新しいCTM議事録テンプレートとして認識できませんでした。\n"
                "添付の「会議議事録テンプレのコピー.xlsx」を選択してください。"
            )

    def _write_basic_info(self, sheet: Worksheet, minutes: MeetingMinutes) -> None:
        sheet["C4"] = build_display_meeting_name(minutes.meeting_name, minutes.short_meeting_name)
        sheet["C5"] = normalize_basic_datetime(minutes.date_time)
        sheet["C6"] = normalize_basic_place(minutes.place)
        sheet["C7"] = minutes.attendees
        self._missing_ranges.extend(["C4:C4", "C5:C5", "C6:C6", "C7:C7"])

    def _write_summary(self, sheet: Worksheet, minutes: MeetingMinutes) -> None:
        value = self._join_items(minutes.summary)
        sheet.cell(row=self.SUMMARY_BODY_ROW, column=2).value = value
        self._set_alignment(sheet.cell(row=self.SUMMARY_BODY_ROW, column=2), horizontal="left", vertical="top")
        self._missing_ranges.append("B10:E12")

    def _rebuild_dynamic_area(self, sheet: Worksheet, minutes: MeetingMinutes) -> None:
        decision_header = self._capture_rows(sheet, 14, 14)[0]
        decision_body = self._capture_rows(sheet, 15, 19)
        spacer_after_decision = self._capture_rows(sheet, 30, 30)[0]
        content_header = self._capture_rows(sheet, 31, 31)[0]
        content_body = self._capture_rows(sheet, 32, 36)
        spacer_after_content = self._capture_rows(sheet, 47, 47)[0]
        todo_header = self._capture_rows(sheet, 48, 48)[0]
        todo_first = self._capture_rows(sheet, 49, 49)[0]
        todo_body = self._capture_rows(sheet, 50, 50)[0]
        spacer_after_todo = self._capture_rows(sheet, 61, 61)[0]
        next_header = self._capture_rows(sheet, 62, 62)[0]
        next_body = self._capture_rows(sheet, 63, 63)[0]

        self._unmerge_rows(sheet, self.DYNAMIC_START_ROW, self.DYNAMIC_END_ROW)
        sheet.delete_rows(self.DYNAMIC_START_ROW, self.DYNAMIC_END_ROW - self.DYNAMIC_START_ROW + 1)

        agendas = minutes.agendas or [AgendaItem(number=1, title="")]
        todos = minutes.todos or [TodoItem(content="")]
        insert_count = (
            1
            + len(agendas) * self.AGENDA_DECISION_ROWS
            + 1
            + 1
            + len(agendas) * self.AGENDA_CONTENT_ROWS
            + 1
            + 1
            + max(len(todos), self.TODO_MIN_ROWS)
            + 1
            + 2
        )
        sheet.insert_rows(self.DYNAMIC_START_ROW, insert_count)

        row = self.DYNAMIC_START_ROW
        self._apply_row_template(sheet, row, decision_header)
        self._merge_if_unmerged(sheet, row, 3, row, 5)
        row += 1
        for agenda in agendas:
            row = self._write_decision_block(sheet, row, agenda, decision_body)

        self._apply_row_template(sheet, row, spacer_after_decision)
        row += 1

        self._apply_row_template(sheet, row, content_header)
        self._merge_if_unmerged(sheet, row, 3, row, 4)
        row += 1
        for agenda in agendas:
            row = self._write_content_block(sheet, row, agenda, content_body)

        self._apply_row_template(sheet, row, spacer_after_content)
        row += 1

        self._apply_row_template(sheet, row, todo_header)
        row += 1
        previous_agenda = None
        for index, todo in enumerate(todos):
            template = todo_first if index == 0 else todo_body
            self._apply_row_template(sheet, row, template)
            self._write_todo_row(sheet, row, todo, previous_agenda)
            self._missing_ranges.append(f"C{row}:E{row}")
            previous_agenda = todo.agenda_label or previous_agenda
            row += 1

        self._apply_row_template(sheet, row, spacer_after_todo)
        row += 1

        self._apply_row_template(sheet, row, next_header)
        self._merge_if_unmerged(sheet, row, 2, row, 5)
        row += 1
        self._apply_row_template(sheet, row, next_body)
        self._merge_if_unmerged(sheet, row, 2, row, 5)
        self._write_next_meeting(sheet, row, minutes)

    def _write_decision_block(self, sheet: Worksheet, row: int, agenda: AgendaItem, templates: list[dict]) -> int:
        lines = self._fit_lines(agenda.decisions, self.AGENDA_DECISION_ROWS)
        for offset, template in enumerate(templates):
            current_row = row + offset
            self._apply_row_template(sheet, current_row, template)
            self._merge_if_unmerged(sheet, current_row, 3, current_row, 5)
            sheet.cell(row=current_row, column=2).value = (
                f"議題{self._circled_number(agenda.number)}" if offset == 0 else ""
            )
            sheet.cell(row=current_row, column=3).value = lines[offset]
            self._set_alignment(sheet.cell(row=current_row, column=3), horizontal="left", vertical="top")
            self._missing_ranges.append(f"C{current_row}:E{current_row}")
        return row + self.AGENDA_DECISION_ROWS

    def _write_content_block(self, sheet: Worksheet, row: int, agenda: AgendaItem, templates: list[dict]) -> int:
        content = list(agenda.content)
        if agenda.pending:
            content.append("【未決事項・継続検討事項】")
            content.extend(agenda.pending)
        lines = self._fit_lines(content, self.AGENDA_CONTENT_ROWS)
        for offset, template in enumerate(templates):
            current_row = row + offset
            self._apply_row_template(sheet, current_row, template)
            self._merge_if_unmerged(sheet, current_row, 3, current_row, 4)
            sheet.cell(row=current_row, column=2).value = (
                f"議題{self._circled_number(agenda.number)}" if offset == 0 else ""
            )
            sheet.cell(row=current_row, column=3).value = lines[offset]
            sheet.cell(row=current_row, column=5).value = agenda.speaker if offset == 0 else ""
            self._set_alignment(sheet.cell(row=current_row, column=3), horizontal="left", vertical="top")
            self._missing_ranges.append(f"C{current_row}:D{current_row}")
        return row + self.AGENDA_CONTENT_ROWS

    def _write_todo_row(self, sheet: Worksheet, row: int, todo: TodoItem, previous_agenda: str | None) -> None:
        label = todo.agenda_label
        sheet.cell(row=row, column=1).value = False
        sheet.cell(row=row, column=2).value = "" if label == previous_agenda else label
        sheet.cell(row=row, column=3).value = todo.content
        sheet.cell(row=row, column=4).value = todo.assignee
        sheet.cell(row=row, column=5).value = todo.due_date
        for column in (3, 4, 5):
            self._set_alignment(sheet.cell(row=row, column=column), horizontal="left", vertical="top")

    def _write_next_meeting(self, sheet: Worksheet, row: int, minutes: MeetingMinutes) -> None:
        date_time = minutes.next_meeting.date_time
        if minutes.next_meeting.note:
            date_time = f"{date_time}（{minutes.next_meeting.note}）" if date_time else minutes.next_meeting.note
        place = minutes.next_meeting.place
        if minutes.next_meeting.place_note:
            place = f"{place}（{minutes.next_meeting.place_note}）" if place else minutes.next_meeting.place_note

        parts = []
        if date_time:
            parts.append(f"日時：{date_time}")
        if place:
            parts.append(f"場所：{place}")
        sheet.cell(row=row, column=2).value = "　".join(parts)
        self._set_alignment(sheet.cell(row=row, column=2), horizontal="left", vertical="center")
        self._missing_ranges.append(f"B{row}:E{row}")

    def _apply_missing_value_highlights(self, sheet: Worksheet) -> None:
        yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        for cell_range in self._missing_ranges:
            start_cell = cell_range.split(":")[0]
            sheet.conditional_formatting.add(
                cell_range,
                FormulaRule(formula=[f'LEN(TRIM({start_cell}&""))=0'], fill=yellow_fill),
            )

    def _format_for_readability(self, sheet: Worksheet) -> None:
        widths = {
            "A": 4,
            "B": 24,
            "C": 64,
            "D": 18,
            "E": 18,
            "F": 4,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal,
                    vertical=cell.alignment.vertical or "top",
                    text_rotation=cell.alignment.text_rotation,
                    wrap_text=True,
                    shrink_to_fit=False,
                    indent=cell.alignment.indent,
                )

        for row_number in range(1, sheet.max_row + 1):
            height = self._estimate_row_height(sheet, row_number)
            if height:
                sheet.row_dimensions[row_number].height = height

    def _estimate_row_height(self, sheet: Worksheet, row_number: int) -> float | None:
        max_lines = 1
        width_by_column = {
            2: 24,
            3: 64,
            4: 18,
            5: 18,
        }
        for column in range(2, 6):
            cell = sheet.cell(row=row_number, column=column)
            if isinstance(cell, MergedCell):
                continue
            value = cell.value
            if value in (None, ""):
                continue
            text = str(value)
            estimated_width = width_by_column.get(column, 16)
            estimated_width = self._merged_width(sheet, row_number, column, estimated_width, width_by_column)
            line_count = 0
            for line in text.splitlines() or [""]:
                line_count += max(1, ceil(self._display_length(line) / max(estimated_width, 1)))
            max_lines = max(max_lines, line_count)

        if max_lines <= 1:
            return 22
        return min(240, max(28, max_lines * 18))

    def _merged_width(
        self,
        sheet: Worksheet,
        row_number: int,
        column: int,
        default_width: int,
        width_by_column: dict[int, int],
    ) -> int:
        for merged_range in sheet.merged_cells.ranges:
            if (
                merged_range.min_row <= row_number <= merged_range.max_row
                and merged_range.min_col == column
                and merged_range.max_col > column
            ):
                return sum(width_by_column.get(col, 16) for col in range(column, merged_range.max_col + 1))
        return default_width

    def _display_length(self, text: str) -> int:
        length = 0
        for char in text:
            length += 2 if ord(char) > 127 else 1
        return length

    def _capture_rows(self, sheet: Worksheet, start: int, end: int) -> list[dict]:
        rows = []
        for row_number in range(start, end + 1):
            rows.append(
                {
                    "height": sheet.row_dimensions[row_number].height,
                    "cells": [self._capture_cell(sheet.cell(row=row_number, column=col)) for col in range(1, 6)],
                }
            )
        return rows

    def _capture_cell(self, cell) -> dict:
        return {
            "style": copy(cell._style),
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "number_format": cell.number_format,
            "protection": copy(cell.protection),
            "value": None if isinstance(cell, MergedCell) else cell.value,
        }

    def _apply_row_template(self, sheet: Worksheet, row_number: int, template: dict) -> None:
        sheet.row_dimensions[row_number].height = template["height"]
        for column, cell_template in enumerate(template["cells"], start=1):
            cell = sheet.cell(row=row_number, column=column)
            if isinstance(cell, MergedCell):
                continue
            cell._style = copy(cell_template["style"])
            cell.font = copy(cell_template["font"])
            cell.fill = copy(cell_template["fill"])
            cell.border = copy(cell_template["border"])
            cell.alignment = copy(cell_template["alignment"])
            cell.number_format = cell_template["number_format"]
            cell.protection = copy(cell_template["protection"])
            cell.value = cell_template["value"]

    def _unmerge_rows(self, sheet: Worksheet, start_row: int, end_row: int) -> None:
        for merged_range in list(sheet.merged_cells.ranges):
            if start_row <= merged_range.min_row and merged_range.max_row <= end_row:
                sheet.unmerge_cells(str(merged_range))

    def _merge_if_unmerged(
        self,
        sheet: Worksheet,
        start_row: int,
        start_column: int,
        end_row: int,
        end_column: int,
    ) -> None:
        target = f"{sheet.cell(start_row, start_column).coordinate}:{sheet.cell(end_row, end_column).coordinate}"
        if target not in {str(range_) for range_ in sheet.merged_cells.ranges}:
            sheet.merge_cells(start_row=start_row, start_column=start_column, end_row=end_row, end_column=end_column)

    def _set_alignment(self, cell, horizontal: str, vertical: str) -> None:
        cell.alignment = Alignment(
            horizontal=horizontal,
            vertical=vertical,
            text_rotation=cell.alignment.text_rotation,
            wrap_text=True,
            shrink_to_fit=False,
            indent=cell.alignment.indent,
        )

    def _fit_lines(self, items: list[str], max_rows: int) -> list[str]:
        if not items:
            return [""] * max_rows
        rows = items[: max_rows - 1]
        rows.append(self._join_items(items[max_rows - 1 :]))
        return (rows + [""] * max_rows)[:max_rows]

    def _join_items(self, items: list[str]) -> str:
        return "\n".join(f"・{item}" for item in items if item)

    def _circled_number(self, number: int) -> str:
        if 1 <= number <= len("①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"):
            return "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"[number - 1]
        return str(number)


def normalize_basic_datetime(value: str, now: datetime | None = None) -> str:
    text = _strip_parenthetical_notes(value)
    date_text = _extract_date_text(text)
    if not date_text:
        current = _jst_now(now)
        date_text = f"{current.year}年{current.month}月{current.day}日"

    time_text = _extract_time_range(text)
    return f"{date_text} {time_text}" if time_text else date_text


def normalize_basic_place(value: str) -> str:
    text = _strip_parenthetical_notes(value)
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"\s+", " ", text).strip(" 　、,。.")
    if not text or text in {"不明", "未定", "なし"}:
        return ""
    if re.search(r"(Zoom|Teams|Google\s*Meet|Meet|Web会議|WEB会議|オンライン|リモート|ウェブ会議)", text, re.IGNORECASE):
        return "オンライン"
    return text


def _strip_parenthetical_notes(value: str) -> str:
    text = value or ""
    previous = None
    while previous != text:
        previous = text
        text = re.sub(r"[（(][^（）()]*[）)]", "", text)
    return text.strip()


def _extract_date_text(text: str) -> str:
    match = re.search(r"(\d{4})年\s*(\d{1,2})月\s*(\d{1,2})日", text)
    if match:
        return f"{int(match.group(1))}年{int(match.group(2))}月{int(match.group(3))}日"
    return ""


def _extract_time_range(text: str) -> str:
    match = re.search(
        r"(\d{1,2})\s*[:：]\s*(\d{2})"
        r"(?:\s*[～〜-]\s*(\d{1,2})\s*[:：]\s*(\d{2}))?",
        text,
    )
    if not match:
        return ""
    start = f"{int(match.group(1))}:{match.group(2)}"
    if not match.group(3):
        return start
    end = f"{int(match.group(3))}:{match.group(4)}"
    return f"{start}〜{end}"


def _jst_now(now: datetime | None) -> datetime:
    if now is None:
        return datetime.now(JST)
    if now.tzinfo is None:
        return now.replace(tzinfo=JST)
    return now.astimezone(JST)
